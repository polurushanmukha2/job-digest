#!/usr/bin/env python3
"""
Daily DevOps / SRE / Platform job scanner.

Sources:
- Greenhouse
- Lever
- Ashby
- Workday direct CXS
- Optional Workday via Apify

Outputs:
- results/jobs_<timestamp>.xlsx
- results/jobs_<timestamp>.md
- results/seen.json

Requires:
- openpyxl==3.1.5
"""

import hashlib
import html
import json
import os
import re
import time
import urllib.parse
import urllib.request
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


RESULTS_DIR = Path("results")
SEEN_FILE = RESULTS_DIR / "seen.json"

GREENHOUSE_CONFIG_URL = (
    "https://raw.githubusercontent.com/ambicuity/New-Grad-Jobs/main/config.yml"
)

PRIORITY_GREENHOUSE = [
    "anthropic", "databricks", "datadog", "elastic", "mongodb",
    "cockroachlabs", "cloudflare", "coinbase", "gitlab", "grafanalabs",
    "pagerduty", "okta", "reddit", "samsara", "stripe", "twilio",
]

LEVER_SLUGS = ["ro"]

ASHBY_SLUGS = [
    "openai", "snowflake", "confluent", "notion", "perplexity",
    "anyscale", "together-ai", "modal", "coreweave", "runway",
    "cursor", "sourcegraph", "ramp", "cohere", "plaid", "supabase",
    "posthog", "huggingface", "vercel", "replit", "retool",
    "robinhood", "chime", "affirm", "scale-ai",
]

WORKDAY_DIRECT = [
    {"name": "Adobe", "url": "https://adobe.wd5.myworkdayjobs.com/external_experienced/jobs"},
    {"name": "Autodesk", "url": "https://autodesk.wd1.myworkdayjobs.com/Ext/jobs"},
    {"name": "Mastercard", "url": "https://mastercard.wd1.myworkdayjobs.com/CorporateCareers/jobs"},
    {"name": "PayPal", "url": "https://paypal.wd1.myworkdayjobs.com/jobs/jobs"},
    {"name": "CVS Health", "url": "https://cvshealth.wd1.myworkdayjobs.com/CVS_Health_Careers/jobs"},
]

# Add blocked Workday career-board URLs here after testing them.
WORKDAY_APIFY = [
    # {"name": "Example", "url": "https://example.wd5.myworkdayjobs.com/External/jobs"},
]

TITLE_INCLUDE = [
    "devops engineer", "site reliability engineer", "site reliability",
    "platform engineer", "cloud engineer", "cloud platform engineer",
    "infrastructure engineer", "production engineer", "reliability engineer",
    "kubernetes engineer", "observability engineer",
    "developer productivity engineer", "developer experience engineer",
    "release engineer", "build engineer", "cloud support engineer",
    "technical support engineer", "devsecops engineer", "systems engineer",
    "software engineer, infrastructure", "software engineer - infrastructure",
]

TITLE_EXCLUDE = [
    "intern", "internship", "new grad", "graduate", "apprentice",
    "manager", "director", "vice president", "vp ", "head of",
    "sales", "marketing", "recruiter", "talent acquisition",
    "account executive", "customer success manager", "product manager",
    "program manager", "project manager", "frontend", "front-end",
    "mobile engineer", "ios engineer", "android engineer",
]

TECH_KEYWORDS = [
    "kubernetes", "terraform", "helm", "argocd", "devops",
    "site reliability", "platform engineering", "infrastructure as code",
    "github actions", "gitlab ci", "jenkins", "aws", "azure", "gcp",
    "linux", "prometheus", "grafana", "opentelemetry", "docker", "ansible",
]

SPONSORSHIP_BLOCKED = [
    "unable to sponsor", "no sponsorship", "will not sponsor",
    "cannot sponsor", "must not require sponsorship",
    "without sponsorship now or in the future",
    "without current or future sponsorship",
    "u.s. citizenship required", "us citizenship required",
    "must be a u.s. citizen", "must be a us citizen",
    "active security clearance", "security clearance required",
    "ts/sci", "secret clearance",
]

SPONSORSHIP_REVIEW = [
    "export controlled", "export control", "export authorization",
    "u.s. person", "us person", "itar", "ear regulations",
]

SPONSORSHIP_POSITIVE = [
    "visa sponsorship is available", "sponsorship is available",
    "will sponsor", "we sponsor", "h-1b sponsorship", "h1b sponsorship",
]

NON_US = [
    "canada", "united kingdom", "london", "toronto", "vancouver",
    "india", "bangalore", "bengaluru", "mumbai", "hyderabad",
    "germany", "berlin", "france", "paris", "ireland", "dublin",
    "australia", "sydney", "singapore", "japan", "tokyo",
    "netherlands", "amsterdam", "poland", "romania", "brazil",
    "mexico", "israel", "tel aviv", "spain", "portugal",
]

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 Chrome/125.0.0.0 Safari/537.36"
    ),
    "Accept": "application/json,text/plain,*/*",
}

HEALTH = []


def get_json(url, timeout=25):
    req = urllib.request.Request(url, headers=HEADERS)
    with urllib.request.urlopen(req, timeout=timeout) as response:
        return json.loads(response.read().decode("utf-8"))


def get_text(url, timeout=25):
    req = urllib.request.Request(url, headers=HEADERS)
    with urllib.request.urlopen(req, timeout=timeout) as response:
        return response.read().decode("utf-8")


def post_json(url, payload, headers=None, timeout=40):
    request_headers = {**HEADERS, "Content-Type": "application/json"}
    if headers:
        request_headers.update(headers)
    req = urllib.request.Request(
        url,
        data=json.dumps(payload).encode("utf-8"),
        headers=request_headers,
        method="POST",
    )
    with urllib.request.urlopen(req, timeout=timeout) as response:
        return json.loads(response.read().decode("utf-8"))


def clean_html(value):
    text = html.unescape(value or "")
    text = re.sub(r"<[^>]+>", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def normalize_url(url):
    if not url:
        return ""
    parsed = urllib.parse.urlsplit(url)
    clean_query = urllib.parse.urlencode([
        (k, v) for k, v in urllib.parse.parse_qsl(parsed.query)
        if not k.lower().startswith("utm_") and k.lower() not in {"source", "gh_src"}
    ])
    return urllib.parse.urlunsplit(
        (parsed.scheme, parsed.netloc, parsed.path.rstrip("/"), clean_query, "")
    )


def make_id(ats, company, native_id, url, title):
    raw = native_id or normalize_url(url) or f"{company}|{title}"
    digest = hashlib.sha256(f"{ats}|{company}|{raw}".encode()).hexdigest()[:24]
    return f"{ats.lower().replace(' ', '_')}_{digest}"


def title_matches(title, department="", description=""):
    t = (title or "").lower()
    combined = f"{title} {department} {description}".lower()

    if any(word in t for word in TITLE_EXCLUDE):
        return False

    if any(word in t for word in TITLE_INCLUDE):
        return True

    if any(x in t for x in ["software engineer", "systems engineer", "support engineer"]):
        return sum(word in combined for word in TECH_KEYWORDS) >= 2

    return False


def location_status(location, description=""):
    text = f"{location} {description}".lower()

    if any(place in text for place in NON_US):
        if not any(us in text for us in ["united states", "usa", "u.s."]):
            return "non_us"

    if "remote" in text:
        return "us_remote"

    return "us_or_unknown"


def sponsorship_status(description):
    text = (description or "").lower()

    for phrase in SPONSORSHIP_POSITIVE:
        if phrase in text:
            return "eligible", phrase

    for phrase in SPONSORSHIP_BLOCKED:
        if phrase in text:
            return "blocked", phrase

    for phrase in SPONSORSHIP_REVIEW:
        if phrase in text:
            return "review_required", phrase

    return "unknown", ""


def build_job(ats, company, native_id, title, location, department, url,
              posted="", description="", employment_type=""):
    description = clean_html(description)
    status, reason = sponsorship_status(description)

    return {
        "id": make_id(ats, company, str(native_id), url, title),
        "title": (title or "").strip(),
        "company": (company or "").strip(),
        "location": (location or "").strip(),
        "department": (department or "").strip(),
        "employment_type": (employment_type or "").strip(),
        "ats": ats,
        "posted": str(posted or "").strip(),
        "url": normalize_url(url),
        "description": description,
        "location_status": location_status(location, description),
        "sponsorship_status": status,
        "sponsorship_reason": reason,
    }


def route(job):
    if not title_matches(job["title"], job["department"], job["description"]):
        return "skipped", "Title/technology mismatch"
    if job["location_status"] == "non_us":
        return "skipped", "Non-US location"
    if job["sponsorship_status"] == "blocked":
        return "blocked", f"Sponsorship restriction: {job['sponsorship_reason']}"
    if job["sponsorship_status"] == "review_required":
        return "review", f"Manual review: {job['sponsorship_reason']}"
    return "matched", ""


def add_routed(job, buckets):
    destination, reason = route(job)
    if reason:
        job["skip_reason"] = reason
    buckets[destination].append(job)


def greenhouse_slugs():
    try:
        text = get_text(GREENHOUSE_CONFIG_URL)
        found = re.findall(r"boards-api\.greenhouse\.io/v1/boards/([^/\"']+)", text)
        return sorted(set(found + PRIORITY_GREENHOUSE))
    except Exception:
        return sorted(set(PRIORITY_GREENHOUSE))


def scrape_greenhouse(slug):
    start = time.time()
    buckets = {"matched": [], "review": [], "blocked": [], "skipped": []}
    try:
        data = get_json(f"https://boards-api.greenhouse.io/v1/boards/{slug}/jobs?content=true")
        jobs = data.get("jobs", [])
        for raw in jobs:
            departments = ", ".join(
                x.get("name", "") for x in raw.get("departments", []) if x.get("name")
            )
            job = build_job(
                "Greenhouse",
                slug.replace("-", " ").title(),
                raw.get("id", ""),
                raw.get("title", ""),
                (raw.get("location") or {}).get("name", ""),
                departments,
                raw.get("absolute_url", ""),
                (raw.get("updated_at", "") or "")[:10],
                raw.get("content", ""),
            )
            add_routed(job, buckets)
        HEALTH.append(["Greenhouse", slug, "ok", len(jobs), len(buckets["matched"]),
                       len(buckets["blocked"]), len(buckets["skipped"]),
                       round(time.time() - start, 2), ""])
    except Exception as exc:
        HEALTH.append(["Greenhouse", slug, "error", 0, 0, 0, 0,
                       round(time.time() - start, 2), str(exc)])
    return buckets


def scrape_lever(slug):
    start = time.time()
    buckets = {"matched": [], "review": [], "blocked": [], "skipped": []}
    try:
        data = get_json(f"https://api.lever.co/v0/postings/{slug}?mode=json")
        jobs = data if isinstance(data, list) else []
        for raw in jobs:
            categories = raw.get("categories", {}) or {}
            location = categories.get("location", "") or categories.get("allLocations", "")
            if isinstance(location, list):
                location = ", ".join(location)
            description = " ".join([
                raw.get("descriptionPlain", "") or "",
                raw.get("additionalPlain", "") or "",
            ])
            posted = ""
            if raw.get("createdAt"):
                posted = datetime.fromtimestamp(
                    raw["createdAt"] / 1000, tz=timezone.utc
                ).strftime("%Y-%m-%d")
            job = build_job(
                "Lever", slug.replace("-", " ").title(), raw.get("id", ""),
                raw.get("text", ""), location,
                categories.get("team", "") or categories.get("department", ""),
                raw.get("hostedUrl", ""), posted, description,
                categories.get("commitment", ""),
            )
            add_routed(job, buckets)
        HEALTH.append(["Lever", slug, "ok", len(jobs), len(buckets["matched"]),
                       len(buckets["blocked"]), len(buckets["skipped"]),
                       round(time.time() - start, 2), ""])
    except Exception as exc:
        HEALTH.append(["Lever", slug, "error", 0, 0, 0, 0,
                       round(time.time() - start, 2), str(exc)])
    return buckets


def scrape_ashby(slug):
    start = time.time()
    buckets = {"matched": [], "review": [], "blocked": [], "skipped": []}
    payload = {
        "operationName": "ApiJobBoardWithTeams",
        "variables": {"organizationHostedJobsPageName": slug},
        "query": """
        query ApiJobBoardWithTeams($organizationHostedJobsPageName: String!) {
          jobBoard: jobBoardWithTeams(
            organizationHostedJobsPageName: $organizationHostedJobsPageName
          ) {
            jobPostings {
              id title locationName jobUrl isRemote publishedDate team { name }
            }
          }
        }
        """,
    }
    try:
        data = post_json(
            "https://jobs.ashbyhq.com/api/non-user-graphql?op=ApiJobBoardWithTeams",
            payload,
            {
                "Origin": "https://jobs.ashbyhq.com",
                "Referer": f"https://jobs.ashbyhq.com/{slug}",
            },
        )
        jobs = data.get("data", {}).get("jobBoard", {}).get("jobPostings", []) or []
        for raw in jobs:
            location = raw.get("locationName", "") or ("Remote" if raw.get("isRemote") else "")
            job = build_job(
                "Ashby", slug.replace("-", " ").title(), raw.get("id", ""),
                raw.get("title", ""), location,
                (raw.get("team") or {}).get("name", ""),
                raw.get("jobUrl", ""), (raw.get("publishedDate", "") or "")[:10],
            )
            add_routed(job, buckets)
        status = "ok" if jobs else "empty_or_blocked"
        HEALTH.append(["Ashby", slug, status, len(jobs), len(buckets["matched"]),
                       len(buckets["blocked"]), len(buckets["skipped"]),
                       round(time.time() - start, 2),
                       "" if jobs else "Zero postings returned"])
    except Exception as exc:
        HEALTH.append(["Ashby", slug, "error", 0, 0, 0, 0,
                       round(time.time() - start, 2), str(exc)])
    return buckets


def workday_parts(base_url):
    host = base_url.split("//", 1)[1].split("/", 1)[0]
    tenant = host.split(".")[0]
    wd = host.split(".")[1]
    site = base_url.split(".myworkdayjobs.com/", 1)[1].split("/", 1)[0]
    cxs = f"https://{tenant}.{wd}.myworkdayjobs.com/wday/cxs/{tenant}/{site}/jobs"
    referer = f"https://{tenant}.{wd}.myworkdayjobs.com/en-US/{site}"
    return tenant, wd, site, cxs, referer


def scrape_workday_direct(company):
    start = time.time()
    buckets = {"matched": [], "review": [], "blocked": [], "skipped": []}
    try:
        tenant, wd, site, cxs, referer = workday_parts(company["url"])
        offset, page_size = 0, 20
        jobs = []
        while True:
            data = post_json(
                cxs,
                {"appliedFacets": {}, "limit": page_size,
                 "offset": offset, "searchText": " "},
                {"Referer": referer, "Accept-Language": "en-US"},
            )
            postings = data.get("jobPostings", []) or []
            total = int(data.get("total", 0) or 0)
            if not postings:
                break
            jobs.extend(postings)
            offset += page_size
            if offset >= total:
                break
            time.sleep(0.25)

        for raw in jobs:
            path = raw.get("externalPath", "") or ""
            url = path
            if path and not path.startswith("http"):
                url = f"https://{tenant}.{wd}.myworkdayjobs.com/{site}{path}"
            job = build_job(
                "Workday", company["name"], path,
                raw.get("title", ""), raw.get("locationsText", "") or "",
                "", url, raw.get("postedOn", "") or "",
            )
            add_routed(job, buckets)

        HEALTH.append(["Workday Direct", company["name"], "ok", len(jobs),
                       len(buckets["matched"]), len(buckets["blocked"]),
                       len(buckets["skipped"]), round(time.time() - start, 2), ""])
    except Exception as exc:
        HEALTH.append(["Workday Direct", company["name"], "error", 0, 0, 0, 0,
                       round(time.time() - start, 2), str(exc)])
    return buckets


def scrape_workday_apify(company):
    start = time.time()
    buckets = {"matched": [], "review": [], "blocked": [], "skipped": []}
    token = os.environ.get("APIFY_TOKEN", "")
    if not token:
        HEALTH.append(["Workday Apify", company["name"], "skipped", 0, 0, 0, 0,
                       0, "APIFY_TOKEN missing"])
        return buckets

    actor = os.environ.get(
        "APIFY_ACTOR_ID",
        "automation-lab~workday-jobs-scraper",
    )
    try:
        run_url = (
            f"https://api.apify.com/v2/acts/{urllib.parse.quote(actor, safe='~')}/runs"
            f"?token={urllib.parse.quote(token)}&waitForFinish=120"
        )
        run = post_json(
            run_url,
            {
                "companyUrl": company["url"].replace("/jobs", ""),
                "searchQuery": "engineer",
                "maxJobs": 250,
                "includeDescription": True,
            },
            timeout=150,
        )
        dataset_id = run.get("data", {}).get("defaultDatasetId", "")
        if not dataset_id:
            raise RuntimeError("No Apify dataset ID returned")

        items = get_json(
            f"https://api.apify.com/v2/datasets/{dataset_id}/items"
            f"?token={urllib.parse.quote(token)}&format=json&clean=true",
            timeout=60,
        )
        jobs = items if isinstance(items, list) else []

        for raw in jobs:
            job = build_job(
                "Workday/Apify", company["name"],
                raw.get("jobId", "") or raw.get("url", ""),
                raw.get("title", "") or raw.get("positionTitle", ""),
                raw.get("location", "") or raw.get("locationsText", ""),
                raw.get("department", "") or "",
                raw.get("url", "") or raw.get("applyUrl", ""),
                raw.get("postedDate", "") or raw.get("postedOn", ""),
                raw.get("description", "") or raw.get("jobDescription", ""),
                raw.get("employmentType", "") or "",
            )
            add_routed(job, buckets)

        HEALTH.append(["Workday Apify", company["name"], "ok" if jobs else "empty",
                       len(jobs), len(buckets["matched"]), len(buckets["blocked"]),
                       len(buckets["skipped"]), round(time.time() - start, 2),
                       "" if jobs else "Zero jobs returned"])
    except Exception as exc:
        HEALTH.append(["Workday Apify", company["name"], "error", 0, 0, 0, 0,
                       round(time.time() - start, 2), str(exc)])
    return buckets


def load_seen():
    if not SEEN_FILE.exists():
        return {}
    try:
        data = json.loads(SEEN_FILE.read_text(encoding="utf-8"))
        return data if isinstance(data, dict) else {}
    except Exception:
        return {}


def save_seen(seen):
    RESULTS_DIR.mkdir(exist_ok=True)
    SEEN_FILE.write_text(json.dumps(dict(sorted(seen.items())), indent=2), encoding="utf-8")


def dedupe(jobs):
    return list({job["id"]: job for job in jobs}.values())


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
GREEN_FILL = PatternFill("solid", fgColor="E2F0D9")
YELLOW_FILL = PatternFill("solid", fgColor="FFF2CC")
RED_FILL = PatternFill("solid", fgColor="FCE4D6")


def format_sheet(ws):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for column in ws.columns:
        width = min(max(len(str(c.value or "")) for c in column) + 2, 55)
        ws.column_dimensions[get_column_letter(column[0].column)].width = max(width, 12)
        for cell in column:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def write_jobs(ws, title, jobs, seen, fill=None):
    ws.title = title
    ws.append([
        "New", "Title", "Company", "Location", "Department", "ATS",
        "Posted", "Sponsorship", "Sponsorship Detail", "First Seen", "Apply",
    ])
    for job in jobs:
        ws.append([
            "YES" if job.get("is_new") else "",
            job["title"], job["company"], job["location"],
            job["department"], job["ats"], job["posted"],
            job["sponsorship_status"], job["sponsorship_reason"],
            seen.get(job["id"], ""), "Apply",
        ])
        row = ws.max_row
        link = ws.cell(row=row, column=11)
        if job["url"]:
            link.hyperlink = job["url"]
            link.style = "Hyperlink"
        if fill:
            for cell in ws[row]:
                cell.fill = fill
        elif job.get("is_new"):
            for cell in ws[row]:
                cell.fill = GREEN_FILL
    format_sheet(ws)


def write_skipped(ws, jobs):
    ws.title = "Skipped Roles"
    ws.append(["Title", "Company", "Location", "Department", "ATS", "Reason", "Apply"])
    for job in jobs:
        ws.append([
            job["title"], job["company"], job["location"],
            job["department"], job["ats"], job.get("skip_reason", ""), "Apply",
        ])
        link = ws.cell(row=ws.max_row, column=7)
        if job["url"]:
            link.hyperlink = job["url"]
            link.style = "Hyperlink"
    format_sheet(ws)


def write_health(ws):
    ws.title = "Scanner Health"
    ws.append([
        "Source", "Company", "Status", "Jobs Returned", "Matched",
        "Blocked", "Skipped", "Duration Seconds", "Message",
    ])
    for row in HEALTH:
        ws.append(row)
    format_sheet(ws)


def write_counts(ws, title, jobs, field):
    ws.title = title
    ws.append([field.title(), "Role Count"])
    counts = {}
    for job in jobs:
        key = job.get(field, "") or "Unknown"
        counts[key] = counts.get(key, 0) + 1
    for key, count in sorted(counts.items(), key=lambda x: x[1], reverse=True):
        ws.append([key, count])
    format_sheet(ws)


def write_excel(path, run_time, new_jobs, all_jobs, review, blocked, skipped, seen):
    wb = openpyxl.Workbook()
    write_jobs(wb.active, "New Jobs", new_jobs, seen, GREEN_FILL)
    write_jobs(wb.create_sheet(), "All Open Matches", all_jobs, seen)
    write_jobs(wb.create_sheet(), "Sponsorship Review", review, seen, YELLOW_FILL)
    write_jobs(wb.create_sheet(), "Explicitly Blocked", blocked, seen, RED_FILL)
    write_skipped(wb.create_sheet(), skipped)
    write_health(wb.create_sheet())
    write_counts(wb.create_sheet(), "By Company", all_jobs, "company")
    write_counts(wb.create_sheet(), "By ATS", all_jobs, "ats")
    write_counts(wb.create_sheet(), "By Location", all_jobs, "location")

    ws = wb.create_sheet("Run Summary")
    ws.append(["Metric", "Value"])
    ws.append(["Run Time UTC", run_time])
    ws.append(["New Jobs", len(new_jobs)])
    ws.append(["All Open Matches", len(all_jobs)])
    ws.append(["Sponsorship Review", len(review)])
    ws.append(["Explicitly Blocked", len(blocked)])
    ws.append(["Skipped Roles", len(skipped)])
    ws.append(["Scanner Errors", sum(row[2] == "error" for row in HEALTH)])
    format_sheet(ws)

    wb.save(path)


def write_markdown(path, run_time, new_jobs, review):
    lines = [
        "# DevOps / SRE / Platform Job Digest",
        "",
        f"**Run:** {run_time}",
        f"**New matching jobs:** {len(new_jobs)}",
        f"**Needs sponsorship review:** {len(review)}",
        "",
        "## New Jobs",
        "",
    ]
    if not new_jobs:
        lines.append("No new matching roles found in this run.")
    else:
        for job in new_jobs:
            lines.extend([
                f"### [{job['title']}]({job['url']})",
                f"**{job['company']}** — {job['location'] or 'Location not listed'}  "
                f"| ATS: {job['ats']}  | Posted: {job['posted'] or 'Not listed'}  "
                f"| Sponsorship: {job['sponsorship_status']}",
                "",
            ])

    if review:
        lines.extend(["## Sponsorship / Export-Control Review", ""])
        for job in review:
            lines.append(
                f"- [{job['title']}]({job['url']}) — "
                f"{job['company']} — {job['location']} — "
                f"{job['sponsorship_reason']}"
            )

    path.write_text("\n".join(lines), encoding="utf-8")


def merge(source, target):
    for key in target:
        target[key].extend(source[key])


def main():
    RESULTS_DIR.mkdir(exist_ok=True)
    buckets = {"matched": [], "review": [], "blocked": [], "skipped": []}

    slugs = greenhouse_slugs()
    print(f"Scanning Greenhouse: {len(slugs)} companies")
    for slug in slugs:
        merge(scrape_greenhouse(slug), buckets)
        time.sleep(0.15)

    print(f"Scanning Lever: {len(LEVER_SLUGS)} companies")
    for slug in LEVER_SLUGS:
        merge(scrape_lever(slug), buckets)

    print(f"Scanning Ashby: {len(ASHBY_SLUGS)} companies")
    for slug in ASHBY_SLUGS:
        merge(scrape_ashby(slug), buckets)

    print(f"Scanning Workday Direct: {len(WORKDAY_DIRECT)} companies")
    for company in WORKDAY_DIRECT:
        merge(scrape_workday_direct(company), buckets)

    print(f"Scanning Workday Apify: {len(WORKDAY_APIFY)} companies")
    for company in WORKDAY_APIFY:
        merge(scrape_workday_apify(company), buckets)

    for key in buckets:
        buckets[key] = dedupe(buckets[key])

    all_jobs = sorted(
        buckets["matched"],
        key=lambda job: (job["posted"], job["company"], job["title"]),
        reverse=True,
    )

    run_time = datetime.now(timezone.utc).strftime("%Y-%m-%d_%H%MZ")
    seen = load_seen()
    new_jobs = []

    for job in all_jobs:
        job["is_new"] = job["id"] not in seen
        if job["is_new"]:
            seen[job["id"]] = run_time
            new_jobs.append(job)

    save_seen(seen)

    xlsx_path = RESULTS_DIR / f"jobs_{run_time}.xlsx"
    md_path = RESULTS_DIR / f"jobs_{run_time}.md"

    write_excel(
        xlsx_path, run_time, new_jobs, all_jobs,
        buckets["review"], buckets["blocked"], buckets["skipped"], seen,
    )
    write_markdown(md_path, run_time, new_jobs, buckets["review"])

    print(f"Done. New jobs: {len(new_jobs)}")
    print(f"All open matches: {len(all_jobs)}")
    print(f"Excel: {xlsx_path}")
    print(f"Markdown: {md_path}")


if __name__ == "__main__":
    main()
